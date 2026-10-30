import asyncio
import json
import os
import tempfile
import time
import unittest
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import FastAPI, WebSocket
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.testclient import TestClient

import star5.web as web
from star5 import effort, engines
from star5.runner import build_sources, detect_lang, normalize
from star5.store import Store
from star5.web import drop_event_logs, install
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
from langchain_core.messages import AIMessage


async def fake_model_call(input_tokens, output_tokens):
    """A real LangChain call, made from a child task as GPT Researcher does."""
    message = AIMessage(content="ok", usage_metadata={"input_tokens": input_tokens, "output_tokens": output_tokens,
                                                       "total_tokens": input_tokens + output_tokens})
    model = GenericFakeChatModel(messages=iter([message]))
    await asyncio.create_task(model.ainvoke("hi"))

H = {"X-Star5": "1"}
SCHEMA_BEFORE_BRIEF = """
CREATE TABLE threads (id TEXT PRIMARY KEY, title TEXT NOT NULL, created REAL NOT NULL, updated REAL NOT NULL);
CREATE TABLE turns (id TEXT PRIMARY KEY, thread_id TEXT NOT NULL REFERENCES threads(id) ON DELETE CASCADE,
  seq INTEGER NOT NULL, kind TEXT NOT NULL, question TEXT NOT NULL, mode TEXT NOT NULL, lang TEXT NOT NULL,
  tone TEXT NOT NULL, max_results INTEGER NOT NULL, status TEXT NOT NULL, answer TEXT NOT NULL DEFAULT '',
  sources TEXT NOT NULL DEFAULT '[]', files TEXT NOT NULL DEFAULT '{}', error TEXT NOT NULL DEFAULT '',
  created REAL NOT NULL, started REAL, finished REAL);
"""


class Fakes:
    def __init__(self):
        self.calls = []
        self.chats = []
        self.briefs = []
        self.gate = None
        self.brief_fails = False
        self.media = []
        self.media_gate = None
        self.media_fails = False
        self.geocoded = []
        self.efforts = []
        self.fast_calls = []
        self.timeouts = []

    async def brief(self, question, language, previous, config_path, extra=""):
        self.briefs.append({"question": question, "language": language, "previous": previous, "extra": extra})
        self.efforts.append(("brief", effort.current()))
        if self.brief_fails:
            raise RuntimeError("model unavailable")
        return f"Question: {question}\nDeliverable: shortlist of named candidates\n"

    async def process_media(self, upload, source, out_dir, config_path):
        self.media.append({"name": upload["name"], "kind": upload["kind"], "bytes": Path(source).read_bytes(),
                           "lang": json.loads(Path(config_path).read_text())["LANGUAGE"]})
        if self.media_gate is not None:
            await self.media_gate.wait()
        await fake_model_call(700, 70)
        if self.media_fails:
            raise RuntimeError("cannot decode")
        (Path(out_dir) / "thumb.jpg").write_bytes(b"JPEG")
        return {"text": f"Contenuto di {upload['name']}: un mulino a vento a Wavre.", "summary": "un mulino",
                "meta": {"width": 10}, "thumb": True}

    async def geocode(self, lat, lon, lang):
        self.geocoded.append((lat, lon, lang))
        return "Aisémont, Wavre, Brabant wallon, Belgio"

    async def research(self, task, report_type, tone, max_results, config_path, sink, doc_path=None):
        self.calls.append({"task": task, "type": report_type, "tone": tone, "max": max_results,
                           "lang": json.loads(Path(config_path).read_text())["LANGUAGE"]})
        if doc_path is not None:
            self.calls[-1]["docs"] = {p.name: p.read_text() for p in sorted(Path(doc_path).iterdir())}
        self.efforts.append(("research", await asyncio.create_task(asyncio.sleep(0, effort.current()))))
        await sink.send_json({"type": "logs", "content": "subqueries", "output": "plan",
                              "metadata": ["query uno", "query due"]})
        await sink.send_json({"type": "logs", "content": "running_subquery_research",
                              "output": "\n🔍 Running research for 'query uno'..."})
        await sink.send_json({"type": "logs", "content": "added_source_url",
                              "output": "✅ Added source url to research: https://www.sqlite.org/wal.html\n"})
        await fake_model_call(1000, 200)
        await fake_model_call(500, 100)
        await asyncio.sleep(0.15)
        if self.gate is not None:
            await self.gate.wait()
        await sink.send_json({"type": "report", "output": "# Titolo\n"})
        answer = "# Titolo\nVedi [WAL](https://www.sqlite.org/wal.html) e [blog](https://example.com/x)."
        Path("outputs").mkdir(exist_ok=True)
        Path("outputs/t1.md").write_text(answer)
        return (answer, [{"url": "https://www.sqlite.org/wal.html", "title": "Write-Ahead Logging"}],
                ["https://other.org/read"], {"md": "outputs/t1.md"})

    async def fast(self, question, language, earlier, extra, files_text, location, max_results, config_path, sink):
        self.fast_calls.append({"question": question, "language": language, "earlier": earlier, "extra": extra,
                                "files": files_text, "location": location, "max": max_results})
        self.efforts.append(("fast", effort.current()))
        active = engines.current()
        self.timeouts.append(active.timeout)
        hit = {"url": "https://phimai.example/", "title": "Phimai", "snippet": "", "age": ""}
        active.stats.record("brave", "ok", [hit])
        active.stats.record("bing", "ok", [hit])
        active.stats.record("duckduckgo", "timeouts")
        active.stats.mark_used([hit])
        await sink.send_json({"type": "logs", "content": "subqueries", "output": "plan", "metadata": ["thai Wavre"]})
        await sink.send_json({"type": "logs", "content": "running_subquery_research",
                              "output": "🔍 Running research for 'thai Wavre'..."})
        await sink.send_json({"type": "logs", "content": "added_source_url",
                              "output": "✅ Added source url to research: https://phimai.example/\n"})
        await sink.send_json({"type": "logs", "content": "writing_report", "output": "Writing"})
        await fake_model_call(400, 90)
        answer = "**Phimai** ([sito](https://phimai.example/))."
        return answer, [{"url": "https://phimai.example/", "title": "Phimai"}], ["https://phimai.example/"], {}
    async def chat(self, report, messages, config_path):
        self.chats.append({"report": report, "messages": messages})
        self.efforts.append(("chat", effort.current()))
        await fake_model_call(300, 50)
        return "Risposta dal rapporto.", [{"url": "https://brave.example/r", "title": "R"}]


class Star5Case(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.cwd = os.getcwd()
        os.chdir(self.tmp.name)
        self.addCleanup(os.chdir, self.cwd)
        base = Path(self.tmp.name) / "config.json"
        base.write_text(json.dumps({"LANGUAGE": "italian", "MAX_SEARCH_RESULTS_PER_QUERY": 8}))
        self.fakes = Fakes()
        site = Path(self.tmp.name) / "site"
        site.mkdir()
        (site / "index.html").write_text("OLD UI")

        @asynccontextmanager
        async def upstream_lifespan(app):
            app.mount("/site", StaticFiles(directory=site), name="frontend")
            app.mount("/outputs", StaticFiles(directory="outputs"), name="outputs")
            yield

        Path("outputs").mkdir()
        self.app = FastAPI(lifespan=upstream_lifespan)
        self.app.mount("/static", StaticFiles(directory=site), name="static")

        @self.app.get("/", response_class=HTMLResponse)
        async def upstream():
            return "UPSTREAM"

        @self.app.get("/api/reports")
        async def reports():
            return {"reports": []}

        @self.app.websocket("/ws")
        async def ws(socket: WebSocket):
            await socket.accept()

        import star5.web as web
        self.addCleanup(setattr, web, "OUTPUTS", web.OUTPUTS)
        web.OUTPUTS = (Path(self.tmp.name) / "outputs").resolve()
        self.runner = install(self.app, data_dir=Path(self.tmp.name) / "data", base_config=str(base),
                              research_fn=self.fakes.research, chat_fn=self.fakes.chat, brief_fn=self.fakes.brief,
                              media_fn=self.fakes.process_media, geocode_fn=self.fakes.geocode,
                              fast_fn=self.fakes.fast)
        self.uploads = self.app.state.star5_uploads
        self.data = Path(self.tmp.name) / "data"
        self.runner.tick_every = 0.02
        self.client = TestClient(self.app)
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)

    def events(self, turn_id, after=0):
        out = []
        with self.client.stream("GET", f"/api/s5/turns/{turn_id}/events?after={after}") as response:
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["x-accel-buffering"], "no")
            for line in response.iter_lines():
                if line.startswith("data: "):
                    event = json.loads(line[6:])
                    out.append(event)
                    if event["t"] == "done":
                        break
        return out

    def start(self, question="Spiega in italiano quando usare la modalità WAL di SQLite", **extra):
        response = self.client.post("/api/s5/threads", headers=H, json={"question": question, **extra})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def upload(self, name="mulino.jpg", data=b"x" * 10, mime="image/jpeg", chunk=4, complete=True):
        created = self.client.post("/api/s5/uploads", headers=H,
                                   json={"name": name, "size": len(data), "mime": mime, "ui": "fr"})
        self.assertEqual(created.status_code, 200, created.text)
        upload_id = created.json()["id"]
        for offset in range(0, len(data), chunk):
            part = data[offset:offset + chunk]
            response = self.client.put(f"/api/s5/uploads/{upload_id}?offset={offset}", content=part,
                                       headers={**H, "Content-Type": "application/octet-stream"})
            self.assertEqual(response.status_code, 200, response.text)
        if complete:
            response = self.client.post(f"/api/s5/uploads/{upload_id}/complete", headers=H)
            self.assertEqual(response.status_code, 200, response.text)
        return upload_id

    def wait_upload(self, upload_id, timeout=5):
        deadline = time.time() + timeout
        while time.time() < deadline:
            state = self.client.get(f"/api/s5/uploads/{upload_id}").json()
            if state["state"] in ("ready", "error"):
                return state
            time.sleep(0.05)
        self.fail("upload not processed")


class AttachmentTests(Star5Case):
    def test_chunked_upload_processed_attached_and_used_by_research(self):
        data = bytes(range(256)) * 5
        upload_id = self.upload(data=data, chunk=300)
        state = self.wait_upload(upload_id)
        self.assertEqual((state["state"], state["kind"], state["summary"], state["thumb"]), ("ready", "image", "un mulino", True))
        self.assertEqual(self.fakes.media[0]["bytes"], data)
        self.assertEqual(self.fakes.media[0]["lang"], "french")
        started = self.start("Che mulino è questo e dove si trova?", attachments=[upload_id],
                             location={"lat": 50.71712, "lon": 4.61234, "accuracy": 35.4, "label": "Wavre"})
        events = self.events(started["turn"])
        phases = [e["phase"] for e in events if e["t"] == "phase"]
        self.assertEqual(phases[:2], ["files", "plan"])
        self.assertIn({"t": "files", "done": 1, "total": 1}, events)
        done = events[-1]["turn"]
        self.assertEqual(done["attachments"], [{"id": upload_id, "kind": "image", "name": "mulino.jpg", "size": len(data)}])
        self.assertEqual(done["location"], {"lat": 50.7171, "lon": 4.6123, "accuracy": 35, "label": "Wavre"})
        extra = self.fakes.briefs[-1]["extra"]
        self.assertIn("un mulino a vento a Wavre", extra)
        self.assertIn("Wavre (latitude 50.717, longitude 4.612, accuracy about 35 m)", extra)
        docs = self.fakes.calls[-1]["docs"]
        self.assertEqual(list(docs), ["01-mulino.jpg.txt"])
        self.assertIn("Attached by the user: image \"mulino.jpg\"", docs["01-mulino.jpg.txt"])
        self.assertFalse((self.data / "turn-docs" / started["turn"]).exists(), "turn documents not cleaned up")
        self.assertEqual((done["tokens_in"], done["tokens_out"], done["llm_calls"]), (2200, 370, 3))
        thumb = self.client.get(f"/api/s5/uploads/{upload_id}/thumb")
        self.assertEqual((thumb.status_code, thumb.content), (200, b"JPEG"))
        original = self.client.get(f"/api/s5/uploads/{upload_id}/file")
        self.assertEqual(original.content, data)
        self.assertEqual(original.headers["content-type"], "application/octet-stream")
        self.assertIn("attachment", original.headers["content-disposition"])
        self.assertIn("sandbox", original.headers["content-security-policy"])
        self.assertEqual(self.client.delete(f"/api/s5/uploads/{upload_id}", headers=H).status_code, 409)
        self.client.delete(f"/api/s5/threads/{started['thread']}", headers=H)
        self.assertFalse((self.data / "uploads" / upload_id).exists())
        self.assertEqual(self.client.get(f"/api/s5/uploads/{upload_id}").status_code, 404)

    def test_research_waits_for_attachments_still_processing(self):
        self.fakes.media_gate = asyncio.Event()
        upload_id = self.upload(name="nota.m4a", mime="audio/mp4")
        started = self.start("Riassumi questa nota vocale", attachments=[upload_id])
        time.sleep(0.3)
        self.assertEqual(self.fakes.calls, [], "research started before the attachment was ready")
        self.assertEqual(self.client.get(f"/api/s5/uploads/{upload_id}").json()["state"], "processing")
        self.client.portal.call(self.fakes.media_gate.set)
        done = self.events(started["turn"])[-1]["turn"]
        self.assertEqual(done["status"], "done")
        self.assertEqual(len(self.fakes.calls), 1)

    def test_failed_attachment_does_not_block_and_is_reported(self):
        self.fakes.media_fails = True
        self.fakes.media_gate = asyncio.Event()
        upload_id = self.upload(name="rotto.mp4", mime="video/mp4")
        started = self.start("Cosa mostra questo video?", attachments=[upload_id])
        self.client.portal.call(self.fakes.media_gate.set)
        done = self.events(started["turn"])[-1]["turn"]
        self.assertEqual(done["status"], "done")
        self.assertEqual(self.client.get(f"/api/s5/uploads/{upload_id}").json()["state"], "error")
        self.assertIn("could not be read", self.fakes.briefs[-1]["extra"])
        self.assertNotIn("docs", self.fakes.calls[-1])
        late = self.upload(name="rotto2.mp4", mime="video/mp4")
        self.assertEqual(self.wait_upload(late)["state"], "error")
        refused = self.client.post(f"/api/s5/threads/{started['thread']}/turns", headers=H,
                                   json={"question": "e questo?", "attachments": [late]})
        self.assertEqual(refused.status_code, 409, "an upload that already failed must not be attachable")

    def test_chat_follow_up_gets_full_attachment_text_and_location(self):
        started = self.start()
        self.events(started["turn"])
        upload_id = self.upload(name="preventivo.pdf", mime="application/pdf")
        self.wait_upload(upload_id)
        turn = self.client.post(f"/api/s5/threads/{started['thread']}/turns", headers=H, json={
            "question": "È conveniente?", "kind": "chat", "attachments": [upload_id],
            "location": {"lat": 50.7, "lon": 4.6, "accuracy": 20}}).json()
        self.assertEqual(self.events(turn["turn"])[-1]["turn"]["status"], "done")
        message = self.fakes.chats[-1]["messages"][-1]["content"]
        self.assertIn("Contenuto di preventivo.pdf", message)
        self.assertIn("The user's current location: unknown area", message)

    def test_chat_follow_up_sees_files_attached_to_earlier_turns(self):
        upload_id = self.upload(name="regolamento.pdf", mime="application/pdf")
        self.wait_upload(upload_id)
        started = self.start("Riassumi il regolamento allegato", attachments=[upload_id])
        self.events(started["turn"])
        turn = self.client.post(f"/api/s5/threads/{started['thread']}/turns", headers=H, json={
            "question": "Cosa dice esattamente l'articolo 3?", "kind": "chat"}).json()
        self.assertEqual(self.events(turn["turn"])[-1]["turn"]["status"], "done")
        message = self.fakes.chats[-1]["messages"][-1]["content"]
        self.assertIn("Files attached earlier in this conversation", message)
        self.assertIn("Contenuto di regolamento.pdf", message)
        self.assertNotIn("Files attached to this message", message)

    def test_upload_validation(self):
        def create(**body):
            return self.client.post("/api/s5/uploads", headers=H, json={"name": "a.jpg", "size": 10, "mime": "", **body})
        self.assertEqual(create(name="virus.exe").status_code, 415)
        self.assertEqual(create(name="disegno.svg").status_code, 415)
        self.assertEqual(create(size=0).status_code, 400)
        self.assertEqual(create(size="10").status_code, 400)
        too_big = create(size=41 * 1024 * 1024)
        self.assertEqual((too_big.status_code, too_big.json()["detail"]), (413, "too_large:40"))
        self.assertEqual(self.client.post("/api/s5/uploads", json={"name": "a.jpg", "size": 1}).status_code, 403)
        self.assertEqual(create(name="registrazione.webm", mime="audio/webm").json()["kind"], "audio")
        self.assertEqual(create(name="clip.webm", mime="video/webm").json()["kind"], "video")
        upload_id = create().json()["id"]
        put = lambda offset, data: self.client.put(f"/api/s5/uploads/{upload_id}?offset={offset}", content=data,
                                                   headers={**H, "Content-Type": "application/octet-stream"})
        self.assertEqual(put(0, b"12345").status_code, 200)
        self.assertEqual(put(0, b"12345").status_code, 409)
        self.assertEqual(put(5, b"123456").status_code, 409)
        self.assertEqual(self.client.post(f"/api/s5/uploads/{upload_id}/complete", headers=H).status_code, 409)
        self.assertEqual(self.client.put(f"/api/s5/uploads/{upload_id}?offset=5", content=b"x" * (4 * 1024 * 1024 + 1),
                                         headers=H).status_code, 413)
        bad = self.client.post("/api/s5/threads", headers=H, json={"question": "q", "attachments": [upload_id]})
        self.assertEqual(bad.status_code, 409, "an upload still uploading must not be attachable")
        self.assertEqual(self.client.get("/api/s5/threads").json()["threads"], [], "failed thread left behind")
        for attachments in (["../etc"], "x", [upload_id, upload_id], ["0" * 16] * 11):
            self.assertEqual(self.client.post("/api/s5/threads", headers=H,
                                              json={"question": "q", "attachments": attachments}).status_code, 400)
        self.assertEqual(self.client.delete(f"/api/s5/uploads/{upload_id}", headers=H).status_code, 200)
        self.assertFalse((self.data / "uploads" / upload_id).exists())

    def test_attachment_cannot_be_reused(self):
        upload_id = self.upload()
        self.wait_upload(upload_id)
        first = self.start(attachments=[upload_id])
        self.events(first["turn"])
        again = self.client.post("/api/s5/threads", headers=H, json={"question": "q", "attachments": [upload_id]})
        self.assertEqual(again.status_code, 409)

    def test_geocode_and_location_validation(self):
        response = self.client.post("/api/s5/geocode", headers=H, json={"lat": 50.71712, "lon": 4.61234, "accuracy": 12.6, "ui": "it"})
        self.assertEqual(response.json(), {"lat": 50.7171, "lon": 4.6123, "accuracy": 13,
                                           "label": "Aisémont, Wavre, Brabant wallon, Belgio"})
        self.assertEqual(self.fakes.geocoded, [(50.7171, 4.6123, "it")])
        for body in ({"lat": 91, "lon": 0}, {"lat": "50", "lon": 4}, {"lat": 50, "lon": 4, "accuracy": -1}, []):
            self.assertEqual(self.client.post("/api/s5/geocode", headers=H, json=body).status_code, 400, body)
            if isinstance(body, dict):
                self.assertEqual(self.client.post("/api/s5/threads", headers=H,
                                                  json={"question": "q", "location": body}).status_code, 400)

    def test_sweep_removes_stale_unattached_uploads_and_orphans(self):
        upload_id = self.upload(complete=False)
        orphan = self.data / "uploads" / "0123456789abcdef"
        orphan.mkdir()
        self.uploads.clock = lambda: time.time() + 25 * 3600
        self.assertEqual(self.uploads.sweep(), 2)
        self.assertFalse((self.data / "uploads" / upload_id).exists())
        self.assertFalse(orphan.exists())
        self.assertEqual(self.client.get(f"/api/s5/uploads/{upload_id}").status_code, 404)


class MediaTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: __import__("shutil").rmtree(self.tmp, ignore_errors=True))
        self.prompts = []

    async def model(self, config_path, prompt, images=()):
        self.prompts.append((prompt, [Path(p).name for p in images]))
        return "DESCRIZIONE"

    def run_process(self, kind, ext, worker_result, source=None):
        from star5 import media

        async def worker(command, src, out):
            self.assertEqual(command, kind)
            for name in worker_result.get("images", []):
                (out / name).write_bytes(b"JPEG")
            return worker_result

        upload = {"kind": kind, "ext": ext, "lang": "fr", "name": f"f.{ext}"}
        return asyncio.run(media.process(upload, source or self.tmp / f"f.{ext}", self.tmp, "cfg", worker, self.model))

    def test_image_video_audio(self):
        image = self.run_process("image", "jpg", {"images": ["image.jpg"], "thumb": "thumb.jpg", "width": 9})
        self.assertEqual((image["text"], image["thumb"], image["meta"]), ("DESCRIZIONE", True, {"width": 9}))
        self.assertIn("in French", self.prompts[-1][0])
        video = self.run_process("video", "mp4", {"images": ["frame00.jpg", "frame01.jpg"], "thumb": "thumb.jpg",
                                                  "duration": 12, "transcript": "[00:01] bonjour", "language": "fr"})
        self.assertEqual(self.prompts[-1][1], ["frame00.jpg", "frame01.jpg"])
        self.assertIn("bonjour", self.prompts[-1][0])
        self.assertIn("Transcript of the audio:\n[00:01] bonjour", video["text"])
        audio = self.run_process("audio", "m4a", {"transcript": "[00:00] ciao", "duration": 3, "truncated": False})
        self.assertEqual((audio["text"], audio["meta"]["duration"]), ("[00:00] ciao", 3))

    def test_pdf_text_and_scanned_pdf(self):
        import pymupdf
        text_pdf, scanned = self.tmp / "t.pdf", self.tmp / "s.pdf"
        with pymupdf.open() as doc:
            for n in range(2):
                doc.new_page().insert_text((72, 72), f"Pagina {n + 1}: " + "contenuto del preventivo " * 20)
            doc.save(text_pdf)
        with pymupdf.open() as doc:
            page = doc.new_page()
            page.draw_rect(pymupdf.Rect(50, 50, 300, 300), color=(0, 0, 0), fill=(0.5, 0.5, 0.5))
            doc.save(scanned)
        result = self.run_process("document", "pdf", {}, text_pdf)
        self.assertIn("[page 2]\nPagina 2", result["text"])
        self.assertEqual(result["meta"], {"pages": 2})
        self.assertEqual(self.prompts, [])
        result = self.run_process("document", "pdf", {}, scanned)
        self.assertEqual((result["text"], result["meta"]["ocr"]), ("DESCRIZIONE", True))
        self.assertEqual(self.prompts[-1][1], ["page01.jpg"])

    def test_office_epub_html_and_text_documents(self):
        import zipfile
        import docx
        from pptx import Presentation
        from pptx.util import Inches
        document = docx.Document()
        document.add_heading("Verbale", 1)
        document.add_paragraph("Gita al museo il 14 novembre.")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Quota", "18 euro"
        document.save(self.tmp / "v.docx")
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Budget 2027"
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "Totale 42.000"
        deck.save(self.tmp / "d.pptx")
        with zipfile.ZipFile(self.tmp / "s.xlsx", "w") as z:
            z.writestr("xl/workbook.xml", '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
                       '<sheet name="Spese" sheetId="1" r:id="rId1"/></sheets></workbook>')
            z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
                       '2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
            z.writestr("xl/sharedStrings.xml", '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                       '<si><t>Voce</t></si><si><t>Affitto</t></si></sst>')
            z.writestr("xl/worksheets/sheet1.xml", '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
                       '2006/main"><sheetData><row><c t="s"><v>0</v></c><c t="inlineStr"><is><t>Euro</t></is></c></row>'
                       '<row><c t="s"><v>1</v></c><c><v>850</v></c></row></sheetData></worksheet>')
        with zipfile.ZipFile(self.tmp / "b.epub", "w") as z:
            z.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="OEBPS/c.opf"/>'
                       '</rootfiles></container>')
            z.writestr("OEBPS/c.opf", '<package xmlns="http://www.idpf.org/2007/opf"><manifest>'
                       '<item id="a" href="a.xhtml"/><item id="b" href="b.xhtml"/></manifest>'
                       '<spine><itemref idref="b"/><itemref idref="a"/></spine></package>')
            z.writestr("OEBPS/a.xhtml", "<html><body><p>Capitolo due</p></body></html>")
            z.writestr("OEBPS/b.xhtml", "<html><body><p>Capitolo uno</p><script>x()</script></body></html>")
        (self.tmp / "p.html").write_text("<html><style>p{}</style><body><h1>Titolo</h1><p>Testo</p></body></html>")
        (self.tmp / "n.txt").write_bytes("Città d'arte".encode("cp1252"))
        results = {name: self.run_process("document", name.split(".")[1], {}, self.tmp / name)["text"]
                   for name in ("v.docx", "d.pptx", "s.xlsx", "b.epub", "p.html", "n.txt")}
        self.assertEqual(results["v.docx"], "Verbale\nGita al museo il 14 novembre.\nQuota | 18 euro")
        self.assertIn("[slide 1]\nBudget 2027\nTotale 42.000", results["d.pptx"])
        self.assertEqual(results["s.xlsx"], "[sheet Spese]\nVoce | Euro\nAffitto | 850")
        self.assertEqual(results["b.epub"], "Capitolo uno\n\nCapitolo due")
        self.assertEqual(results["p.html"], "Titolo\nTesto")
        self.assertEqual(results["n.txt"], "Città d'arte")
        with self.assertRaises(ValueError):
            self.run_process("document", "exe", {}, self.tmp / "n.txt")

    def test_summary_skips_markdown(self):
        from star5.media import first_line
        self.assertEqual(first_line("**Descrizione dell'immagine**\n\nUna **farmacia** a Bergamo.\n- Orari"),
                         "Una farmacia a Bergamo. Orari")
        self.assertEqual(first_line("Il video mostra un gatto marmorizzato in una foresta tropicale.\n\n"
                                    "**Contenuto visivo**\n- foglie"),
                         "Il video mostra un gatto marmorizzato in una foresta tropicale.")
        self.assertEqual(first_line("## Video\nUn gatto."), "Un gatto.")
        self.assertEqual(first_line("[00:00] ciao file_name"), "[00:00] ciao file_name")

    def test_classify_and_names(self):
        from star5.media import classify, clean_name
        self.assertEqual(classify("IMG_1.HEIC", ""), ("image", "heic"))
        self.assertEqual(classify("a.webm", "audio/webm"), ("audio", "webm"))
        self.assertEqual(classify("a.webm", ""), ("video", "webm"))
        self.assertEqual(classify("a.svg", "image/svg+xml"), (None, "svg"))
        self.assertEqual(clean_name("../../etc/pass\x00wd"), "passwd")

    def test_geocoder_label_cache_and_rate_limit(self):
        from star5.media import Geocoder
        calls, clock = [], [100.0]

        async def fetch(lat, lon, lang):
            calls.append((lat, lon, lang, clock[0]))
            return {"address": {"road": "Rue X", "house_number": "3", "neighbourhood": "Aisémont", "town": "Wavre",
                                "state": "Brabant wallon", "country": "Belgique"}}

        geocoder = Geocoder(fetch, clock=lambda: clock[0])
        label = asyncio.run(geocoder(50.71712, 4.61234, "fr"))
        self.assertEqual(label, "Aisémont, Wavre, Brabant wallon, Belgique")
        self.assertEqual(calls[0][:2], (50.717, 4.612), "coordinates must be rounded before leaving the server")
        asyncio.run(geocoder(50.71749, 4.61201, "fr"))
        self.assertEqual(len(calls), 1, "same area must come from the cache")


class FastModeTests(Star5Case):
    def test_fast_mode_skips_brief_and_agent_and_defaults_to_low_reasoning(self):
        started = self.start("mi trovi un buon ristorante thai a Wavre?", mode="fast",
                             location={"lat": 50.7171, "lon": 4.6123, "accuracy": 30, "label": "Wavre"})
        events = self.events(started["turn"])
        done = events[-1]["turn"]
        self.assertEqual((done["status"], done["mode"], done["effort"]), ("done", "fast", ""))
        self.assertEqual(self.fakes.briefs, [], "Veloce must not write a research brief")
        self.assertEqual(self.fakes.calls, [], "Veloce must not run the GPT Researcher agent")
        self.assertEqual(self.fakes.efforts, [("fast", "low")])
        call = self.fakes.fast_calls[0]
        self.assertEqual((call["language"], call["max"]), ("italian", 8))
        self.assertIn("Wavre (latitude 50.717", call["location"])
        phases = [e["phase"] for e in events if e["t"] == "phase"]
        self.assertEqual(phases, ["plan", "search", "read", "write"])
        self.assertIn({"t": "queries", "items": ["thai Wavre"]}, events)
        self.assertEqual({k: done["sources"][0][k] for k in ("url", "title", "domain", "cited")},
                         {"url": "https://phimai.example/", "title": "Phimai", "domain": "phimai.example", "cited": True})
        self.assertEqual((done["tokens_in"], done["tokens_out"], done["llm_calls"]), (400, 90, 1))
        self.assertEqual(self.fakes.timeouts, [engines.MODE_TIMEOUTS["fast"]])
        self.assertEqual(done["sources"][0]["engines"], ["brave", "bing"])
        self.assertEqual(done["engines"]["brave"], {"used": 1, "read": 1, "cited": 1, "queries": 1, "timeouts": 0, "errors": 0})
        self.assertEqual(done["engines"]["duckduckgo"]["timeouts"], 1)
        self.assertEqual(self.client.get(f"/api/s5/threads/{started['thread']}").json()["turns"][0]["engines"], done["engines"])

    def test_explicit_effort_wins_and_follow_up_web_search_keeps_context(self):
        started = self.start("Prima domanda sulle pompe di calore in Belgio", mode="fast", effort="high")
        self.events(started["turn"])
        self.assertEqual(self.fakes.efforts, [("fast", "high")])
        upload_id = self.upload(name="preventivo.pdf", mime="application/pdf")
        self.wait_upload(upload_id)
        turn = self.client.post(f"/api/s5/threads/{started['thread']}/turns", headers=H, json={
            "question": "E per un appartamento?", "kind": "research", "mode": "fast",
            "attachments": [upload_id]}).json()
        self.assertEqual(self.events(turn["turn"])[-1]["turn"]["status"], "done")
        call = self.fakes.fast_calls[-1]
        self.assertEqual(call["earlier"], "Prima domanda sulle pompe di calore in Belgio")
        self.assertIn("Contenuto di preventivo.pdf", call["files"])
        self.assertIn("preventivo.pdf", call["extra"])

    def test_fast_answer_does_not_wait_behind_a_long_research(self):
        self.fakes.gate = asyncio.Event()
        slow = self.start("Una ricerca lunga che resta in corso", mode="quick")
        time.sleep(0.2)
        quick = self.start("Una domanda veloce mentre l'altra è in corso", mode="fast")
        self.assertEqual(self.events(quick["turn"])[-1]["turn"]["status"], "done")
        self.client.portal.call(self.fakes.gate.set)
        self.assertEqual(self.events(slow["turn"])[-1]["turn"]["status"], "done")


class FastPipelineTests(unittest.TestCase):
    def setUp(self):
        import httpx
        from star5 import fast
        self.fast, self.httpx = fast, httpx
        self.events = []
        self.searched = []

    class Sink:
        def __init__(self, events):
            self.events = events

        async def send_json(self, data):
            self.events.append(data)

    def transport(self, pages):
        async def handler(request):
            body = pages.get(str(request.url))
            if body is None:
                return self.httpx.Response(404)
            if body == "slow":
                await asyncio.sleep(0.3)
                return self.httpx.Response(200, text="<p>too late</p>", headers={"content-type": "text/html"})
            return self.httpx.Response(200, text=body, headers={"content-type": "text/html; charset=utf-8"})
        return self.httpx.MockTransport(handler)

    def fake_search(self, query, count):
        self.searched.append((query, count))
        word = query.split()[0]
        return [{"url": f"https://{word}.example/{i}", "title": f"{query} {i}", "snippet": "Desc & more",
                 "age": "", "engines": ["brave", "bing"] if i == 0 else ["duckduckgo"]} for i in range(3)] + [
                {"url": "https://shared.example/", "title": "Shared", "snippet": "s", "age": "", "engines": ["bing"]}]

    def run_fast(self, **overrides):
        captured = {}

        async def planner(question, earlier, extra, config_path):
            return ["alpha query", "beta query"]

        async def writer(question, language, earlier, files_text, location, results, config_path):
            captured["block"] = self.fast.sources_block(results)
            captured["results"] = results
            return "Risposta [alpha](https://alpha.example/0)."

        async def files_fn(answer, question):
            return {"md": "outputs/x.md", "pdf": ""}

        async def guard(request):
            if request.url.host.endswith(".private"):
                raise self.httpx.RequestError("blocked", request=request)

        kwargs = dict(planner=planner, writer=writer, files_fn=files_fn, guard=guard, search=self.fake_search,
                      transport=self.transport(overrides.pop("pages", {})))
        kwargs.update(overrides)
        out = asyncio.run(self.fast.fast_research("Domanda?", "italian", "", "", "", "", 8, "cfg",
                                                  self.Sink(self.events), **kwargs))
        return out, captured

    def test_search_read_and_answer(self):
        pages = {"https://alpha.example/0": "<html><nav>menu</nav><main><h1>Alpha</h1><p>Contenuto utile alpha</p>"
                                            "<script>x()</script></main></html>",
                 "https://beta.example/0": "<html><body><p>Beta text</p></body></html>"}
        (answer, sources, visited, files), captured = self.run_fast(pages=pages)
        self.assertEqual(answer, "Risposta [alpha](https://alpha.example/0).")
        self.assertEqual(sorted(self.searched), [("alpha query", 8), ("beta query", 8)])
        self.assertEqual([s["url"] for s in sources][:4], ["https://alpha.example/0", "https://beta.example/0",
                                                          "https://alpha.example/1", "https://beta.example/1"])
        self.assertEqual(sum(1 for s in sources if s["url"] == "https://shared.example/"), 1, "duplicates merged")
        self.assertEqual(visited, ["https://alpha.example/0", "https://beta.example/0"])
        self.assertEqual(files, {"md": "outputs/x.md", "pdf": ""})
        self.assertIn("Page text:\nAlpha\nContenuto utile alpha", captured["block"])
        self.assertNotIn("menu", captured["block"])
        self.assertNotIn("x()", captured["block"])
        kinds = [e.get("content") or e["type"] for e in self.events]
        self.assertEqual(kinds, ["subqueries", "running_subquery_research", "running_subquery_research",
                                 "added_source_url", "added_source_url", "writing_report", "report"])

    def test_slow_pages_dropped_at_the_deadline(self):
        old = self.fast.FETCH_DEADLINE
        self.fast.FETCH_DEADLINE = 0.15
        self.addCleanup(setattr, self.fast, "FETCH_DEADLINE", old)
        pages = {"https://alpha.example/0": "slow", "https://beta.example/0": "<p>ok beta</p>"}
        (_, sources, visited, _), _ = self.run_fast(pages=pages)
        self.assertGreaterEqual(len(sources), 7)
        self.assertEqual(visited, ["https://beta.example/0"])

    def test_no_results_fail_clearly(self):
        with self.assertRaisesRegex(RuntimeError, "no results"):
            self.run_fast(search=lambda query, count: [])

    def test_query_helpers(self):
        fast = self.fast
        brief = "Question: Qual è la differenza tra pompa aria-acqua e aria-aria?\n\nGoal: " + "parola " * 200
        self.assertEqual(fast.short_query(brief), "Qual è la differenza tra pompa aria-acqua e aria-aria?")
        self.assertLessEqual(len(fast.short_query("x" * 900)), 380)
        self.assertLessEqual(len(fast.short_query(" ".join(["w"] * 80)).split()), 45)
        self.assertEqual(fast.short_query("  breve   domanda "), "breve domanda")
        self.assertEqual(fast.parse_queries('Ecco: {"queries": ["a b", "A B", "c", "d", "e"]}', "q"), ["a b", "c", "d"])
        self.assertEqual(fast.parse_queries("niente json", "la domanda"), ["la domanda"])
        self.assertEqual(fast.parse_queries('{"queries": [1, null]}', "q"), ["q"])

    def test_interleave_merges_engines_of_the_same_url(self):
        a = [{"url": "https://www.x.example/p/", "engines": ["bing"]}]
        b = [{"url": "https://x.example/p", "engines": ["brave"]}, {"url": "https://y.example/", "engines": ["duckduckgo"]}]
        merged = self.fast.interleave([a, b])
        self.assertEqual([m["url"] for m in merged], ["https://www.x.example/p/", "https://y.example/"])
        self.assertEqual(merged[0]["engines"], ["brave", "bing"])

    def test_only_public_hosts(self):
        is_public = self.fast.is_public
        for host in ("localhost", "10.183.20.1", "192.168.178.150", "127.0.0.1", "169.254.169.254", "[::1]",
                     "printer.local", "fd00::1", ""):
            self.assertFalse(asyncio.run(is_public(host)), host)
        self.assertTrue(asyncio.run(is_public("8.8.8.8")))


class EngineTests(unittest.TestCase):
    def setUp(self):
        from unittest import mock
        from star5 import engines
        self.engines, self.mock = engines, mock
        os.environ["BRAVE_API_KEY"] = "unit-test-brave"
        self.addCleanup(os.environ.pop, "BRAVE_API_KEY", None)

    def fake_engines(self, behaviour):
        def make(name):
            def run(query, count, timeout):
                kind = behaviour[name]
                if kind == "slow":
                    time.sleep(timeout + 0.4)
                    return [{"url": "https://late.example/", "title": "", "snippet": "", "age": ""}]
                if kind == "error":
                    raise RuntimeError("blocked")
                return [{"url": u, "title": name, "snippet": f"{name} snippet", "age": ""} for u in kind]
            return run
        return self.mock.patch.dict(self.engines.ENGINE_FUNCS, {n: make(n) for n in behaviour})

    def test_parallel_engines_timeout_errors_merge_and_stats(self):
        behaviour = {"brave": ["https://a.example/", "https://b.example/"],
                     "bing": ["https://www.a.example", "https://c.example/"], "duckduckgo": "slow"}
        with self.fake_engines(behaviour), self.engines.session("fast") as active:
            self.assertEqual(active.timeout, 0 + self.engines.MODE_TIMEOUTS["fast"])
            started = time.monotonic()
            results = self.engines.search_all("q", 5, timeout=0.3)
            elapsed = time.monotonic() - started
            active.stats.mark_used(results)
            summary = active.stats.summary(["https://a.example/"], ["https://c.example/"])
        self.assertLess(elapsed, 0.6, "the slow engine must not hold the others")
        self.assertEqual([r["url"] for r in results], ["https://a.example/", "https://b.example/", "https://c.example/"])
        self.assertEqual(results[0]["engines"], ["brave", "bing"])
        self.assertEqual(summary["brave"], {"used": 2, "read": 1, "cited": 0, "queries": 1, "timeouts": 0, "errors": 0})
        self.assertEqual(summary["bing"], {"used": 2, "read": 1, "cited": 1, "queries": 1, "timeouts": 0, "errors": 0})
        self.assertEqual(summary["duckduckgo"], {"used": 0, "read": 0, "cited": 0, "queries": 1, "timeouts": 1,
                                                 "errors": 0})
        self.assertIsNone(self.engines.current())

    def test_errors_and_missing_brave_key(self):
        behaviour = {"brave": ["https://a.example/"], "bing": "error", "duckduckgo": ["https://d.example/"]}
        stats = self.engines.Stats()
        with self.fake_engines(behaviour):
            self.engines.search_all("q", 5, timeout=1, stats=stats)
            os.environ.pop("BRAVE_API_KEY")
            results = self.engines.search_all("q", 5, timeout=1, stats=stats)
        self.assertEqual([r["url"] for r in results], ["https://d.example/"])
        self.assertEqual(stats.engines["bing"]["errors"], 2)
        self.assertEqual(stats.engines["brave"]["queries"], 1, "no Brave call without a key")

    def test_mode_timeouts(self):
        timeouts = [self.engines.MODE_TIMEOUTS[m] for m in ("fast", "quick", "deep")]
        self.assertEqual(timeouts, sorted(timeouts))
        self.assertEqual(len(set(timeouts)), 3, "each mode has its own timeout")
        with self.engines.session("deep") as active:
            self.assertEqual(active.timeout, self.engines.MODE_TIMEOUTS["deep"])

    def test_ddgs_engines_drop_ads_and_choose_backends(self):
        calls = []

        class FakeDDGS:
            def __init__(self, timeout=5):
                calls.append({"timeout": timeout})

            def text(self, query, **kwargs):
                calls[-1].update(kwargs)
                return [{"href": "https://www.bing.com/aclick?ld=x", "title": "ad"},
                        {"href": "https://real.example/", "title": "<b>Real</b>", "body": "a &amp; b"},
                        {"href": "ftp://no"}, "junk"]
        with self.mock.patch("ddgs.DDGS", FakeDDGS):
            bing = self.engines.bing("q", 6, 4.5)
            ddg = self.engines.duckduckgo("q", 6, 4.5)
        self.assertEqual(bing, [{"url": "https://real.example/", "title": "Real", "snippet": "a & b", "age": ""}])
        self.assertEqual(ddg, bing)
        self.assertEqual((calls[0]["backend"], calls[0]["region"]), ("bing", "us-en"))
        self.assertIn(calls[0]["timeout"], (3, 4))
        self.assertEqual((calls[1]["backend"], calls[1]["region"]), ("duckduckgo", "wt-wt"))

    def test_ddgs_second_region_when_the_first_fails(self):
        regions = []

        class Flaky:
            def __init__(self, timeout=5):
                pass

            def text(self, query, region, **kwargs):
                regions.append(region)
                if len(regions) == 1:
                    raise RuntimeError("No results found.")
                return [{"href": "https://ok.example/", "title": "ok"}]

        class Down(Flaky):
            def text(self, query, region, **kwargs):
                regions.append(region)
                raise RuntimeError("No results found.")
        with self.mock.patch("ddgs.DDGS", Flaky):
            self.assertEqual(len(self.engines.duckduckgo("q", 5, 6)), 1)
        self.assertEqual(regions, ["wt-wt", "us-en"])
        regions.clear()
        with self.mock.patch("ddgs.DDGS", Down), self.assertRaisesRegex(RuntimeError, "No results"):
            self.engines.bing("q", 5, 6)
        self.assertEqual(regions, ["us-en", "uk-en"])
        regions.clear()

        class SlowDown(Down):
            def text(self, query, region, **kwargs):
                time.sleep(0.2)
                return super().text(query, region, **kwargs)
        with self.mock.patch("ddgs.DDGS", SlowDown), self.assertRaises(RuntimeError):
            self.engines.bing("q", 5, 1.6)
        self.assertEqual(regions, ["us-en"], "no second try without enough time left")

    def test_brave_parses_and_retries_429(self):
        import httpx
        replies = [httpx.Response(429, json={}), httpx.Response(200, json={"web": {"results": [
            {"url": "https://a.example/", "title": "<strong>A</strong>", "description": "d", "extra_snippets": ["e"],
             "age": "2 days"}, {"url": "javascript:x"}]}})]
        with self.mock.patch.object(httpx, "get", side_effect=lambda *a, **k: replies.pop(0)), \
                self.mock.patch.object(self.engines.time, "sleep"):
            found = self.engines.brave("q", 5, 6)
        self.assertEqual(found, [{"url": "https://a.example/", "title": "A", "snippet": "d … e", "age": "2 days"}])
        with self.mock.patch.object(httpx, "get", return_value=httpx.Response(402, json={})):
            with self.assertRaisesRegex(RuntimeError, "402"):
                self.engines.brave("q", 5, 6)

    def test_gpt_researcher_retriever_uses_all_engines(self):
        from gpt_researcher.retrievers.brave.brave import BraveSearch
        self.engines.install_gpt_researcher()
        self.engines.install_gpt_researcher()
        self.assertEqual(BraveSearch("Question: Dove trovo X?\n\nGoal: " + "lungo " * 300).query, "Dove trovo X?")
        behaviour = {"brave": ["https://a.example/"], "bing": ["https://b.example/"],
                     "duckduckgo": [f"https://d{i}.example/" for i in range(10)]}
        with self.fake_engines(behaviour), self.engines.session("quick") as active:
            results = BraveSearch("breve").search(max_results=3)
            summary = active.stats.summary()
        self.assertEqual(len(results), 4, "merged list capped at 1.5 times the per-engine count (minimum 4)")
        self.assertEqual(results[0], {"href": "https://a.example/", "title": "brave", "body": "brave snippet"})
        self.assertEqual(results[1]["body"], "bing snippet"[:100])
        self.assertEqual(summary["duckduckgo"]["used"], 2)

class EffortTests(Star5Case):
    def test_effort_reaches_brief_research_and_nested_tasks_then_follow_up_chat(self):
        config = self.client.get("/api/s5/config").json()
        self.assertEqual(config["efforts"], ["low", "medium", "high", "xhigh", "max"])
        started = self.start(effort="xhigh")
        done = self.events(started["turn"])[-1]["turn"]
        self.assertEqual((done["status"], done["effort"]), ("done", "xhigh"))
        self.assertEqual(self.fakes.efforts, [("brief", "xhigh"), ("research", "xhigh")])
        turn = self.client.post(f"/api/s5/threads/{started['thread']}/turns", headers=H,
                                json={"question": "E quindi?", "kind": "chat", "effort": "low"}).json()
        self.assertEqual(self.events(turn["turn"])[-1]["turn"]["effort"], "low")
        self.assertEqual(self.fakes.efforts[-1], ("chat", "low"))
        self.assertEqual(effort.current(), "", "effort must not leak outside the job")

    def test_default_effort_is_the_models_and_invalid_is_rejected(self):
        started = self.start()
        done = self.events(started["turn"])[-1]["turn"]
        self.assertEqual(done["effort"], "")
        self.assertEqual(self.fakes.efforts, [("brief", ""), ("research", "")])
        for bad in ("ultra", "none", 3):
            response = self.client.post("/api/s5/threads", headers=H, json={"question": "q", "effort": bad})
            self.assertEqual(response.status_code, 400, bad)
        self.assertEqual(len(self.client.get("/api/s5/threads").json()["threads"]), 1)

    def test_hook_sets_reasoning_effort_on_the_real_langchain_client(self):
        from gpt_researcher.llm_provider.generic.base import GenericLLMProvider
        effort.install()
        effort.install()
        args = dict(model="frontier", api_key="unit-test", azure_endpoint="https://proxy.test",
                    api_version="2024-06-01", max_tokens=100)
        self.assertIsNone(GenericLLMProvider.from_provider("azure_openai", **args).llm.reasoning_effort)
        with effort.reasoning("max"):
            llm = GenericLLMProvider.from_provider("azure_openai", **args).llm
            self.assertEqual(llm.reasoning_effort, "max")
            payload = llm._get_request_payload([AIMessage(content="ciao")])
            self.assertEqual(payload["reasoning_effort"], "max")
            self.assertNotIn("reasoning", payload, "must stay a Chat Completions request")
            self.assertEqual(effort.apply("anthropic", {}), {})
        with effort.reasoning("bogus"):
            self.assertEqual(effort.current(), "")


class ApiTests(Star5Case):
    def test_index_replaces_upstream_page_with_strict_csp(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("UPSTREAM", response.text)
        self.assertNotIn("/site/", response.text)
        self.assertIn("script-src 'self'", response.headers["content-security-policy"])
        for asset in ("/s5/app.js", "/s5/vendor/showdown.min.js", "/s5/vendor/purify.min.js"):
            self.assertEqual(self.client.get(asset).status_code, 200, asset)

    def test_app_is_an_installable_pwa(self):
        csp = self.client.get("/").headers["content-security-policy"]
        self.assertIn("manifest-src 'self'", csp)
        self.assertIn("worker-src 'self'", csp)
        worker = self.client.get("/sw.js")
        self.assertEqual(worker.status_code, 200)
        self.assertIn("javascript", worker.headers["content-type"])
        self.assertIn("no-store", worker.headers["cache-control"])
        manifest = self.client.get("/s5/manifest.json").json()
        self.assertEqual((manifest["start_url"], manifest["scope"], manifest["display"]), ("/", "/", "standalone"))
        sizes = {(i["sizes"], i["purpose"]) for i in manifest["icons"]}
        self.assertTrue({("192x192", "any"), ("512x512", "any"), ("512x512", "maskable")} <= sizes)
        for icon in manifest["icons"]:
            self.assertEqual(self.client.get(icon["src"]).status_code, 200, icon["src"])
        paths = [getattr(r, "path", "") for r in self.app.router.routes]
        self.assertEqual(paths.count("/sw.js"), 1)

    def test_old_ui_and_its_api_are_gone_after_startup(self):
        for path in ("/site/index.html", "/static/index.html", "/api/reports", "/ws", "/docs", "/openapi.json"):
            self.assertEqual(self.client.get(path).status_code, 404, path)
        Path("outputs/report.md").write_text("ok")
        self.assertEqual(self.client.get("/outputs/report.md").text, "ok")
        paths = [getattr(r, "path", "") for r in self.app.router.routes]
        self.assertEqual(paths.count("/outputs"), 1)
        self.assertEqual(paths.count("/"), 1)
        self.assertEqual(self.client.get("/api/s5/config").status_code, 200)

    def test_research_runs_in_background_and_streams_progress(self):
        started = self.start(mode="detailed", tone="Critical", max_results=10)
        self.assertEqual(started["lang"], "it")
        events = self.events(started["turn"])
        kinds = [e["t"] for e in events]
        self.assertIn("queries", kinds)
        self.assertIn({"t": "query", "q": "query uno"}, events)
        self.assertIn({"t": "source", "url": "https://www.sqlite.org/wal.html", "domain": "sqlite.org"}, events)
        self.assertIn({"t": "chunk", "text": "# Titolo\n"}, events)
        phases = [e["phase"] for e in events if e["t"] == "phase"]
        self.assertEqual(phases, ["plan", "search", "read", "write"])
        done = events[-1]["turn"]
        self.assertEqual(done["status"], "done")
        self.assertEqual([s["url"] for s in done["sources"]],
                         ["https://www.sqlite.org/wal.html", "https://example.com/x", "https://other.org/read"])
        self.assertEqual(done["sources"][0]["title"], "Write-Ahead Logging")
        self.assertEqual([s["cited"] for s in done["sources"]], [True, True, False])
        self.assertEqual(self.fakes.calls[0], {
            "task": "Question: Spiega in italiano quando usare la modalità WAL di SQLite\n"
                    "Deliverable: shortlist of named candidates",
            "type": "detailed_report", "tone": "Critical", "max": 10, "lang": "italian"})
        self.assertEqual(self.fakes.briefs[0]["language"], "italian")
        self.assertIn({"t": "brief", "text": self.fakes.calls[0]["task"]}, events)
        self.assertEqual(done["brief"], self.fakes.calls[0]["task"])
        self.assertEqual((done["tokens_in"], done["tokens_out"], done["llm_calls"]), (1500, 300, 2))
        live = [e for e in events if e["t"] == "tokens"]
        self.assertTrue(live, "no live token event")
        self.assertEqual(live[-1], {"t": "tokens", "tokens_in": 1500, "tokens_out": 300, "llm_calls": 2})
        thread = self.client.get(f"/api/s5/threads/{started['thread']}").json()
        self.assertEqual(thread["turns"][0]["files"], {"md": "outputs/t1.md"})
        listing = self.client.get("/api/s5/threads").json()["threads"]
        self.assertEqual(listing[0]["status"], "done")

    def test_resume_after_reconnect_and_finished_turn_replays_final_state(self):
        started = self.start()
        events = self.events(started["turn"])
        tail = self.events(started["turn"], after=len(events) - 1)
        self.assertEqual([e["t"] for e in tail], ["done"])
        self.runner.jobs.clear()
        again = self.events(started["turn"])
        self.assertEqual(len(again), 1)
        self.assertEqual(again[0]["turn"]["status"], "done")

    def test_report_language_follows_the_question_or_explicit_choice(self):
        french = self.start("Explique quand utiliser le mode WAL de SQLite et ses limites", ui="it")
        self.assertEqual(french["lang"], "fr")
        self.events(french["turn"])
        self.assertEqual(self.fakes.calls[-1]["lang"], "french")
        forced = self.start("Explique quand utiliser le mode WAL de SQLite et ses limites", lang="en")
        self.assertEqual(forced["lang"], "en")
        short = self.start("WAL?", ui="fr")
        self.assertEqual(short["lang"], "fr")

    def test_follow_up_chat_uses_report_and_previous_chat(self):
        started = self.start()
        self.events(started["turn"])
        first = self.client.post(f"/api/s5/threads/{started['thread']}/turns", headers=H,
                                 json={"question": "E con NFS?", "kind": "chat"}).json()
        self.assertEqual(self.events(first["turn"])[-1]["turn"]["answer"], "Risposta dal rapporto.")
        second = self.client.post(f"/api/s5/threads/{started['thread']}/turns", headers=H,
                                  json={"question": "Riassumi", "kind": "chat"}).json()
        done = self.events(second["turn"])[-1]["turn"]
        self.assertEqual(done["mode"], "chat")
        self.assertEqual(done["sources"][0]["url"], "https://brave.example/r")
        last = self.fakes.chats[-1]
        self.assertIn("sqlite.org/wal.html", last["report"])
        self.assertEqual([m["role"] for m in last["messages"]], ["user", "assistant", "user"])
        self.assertIn("Riassumi", last["messages"][-1]["content"])
        self.assertEqual((done["tokens_in"], done["tokens_out"], done["llm_calls"]), (300, 50, 1))
        thread = self.client.get(f"/api/s5/threads/{started['thread']}").json()
        self.assertEqual((thread["tokens_in"], thread["tokens_out"], thread["llm_calls"]), (2100, 400, 4))

    def test_follow_up_research_carries_previous_question(self):
        started = self.start()
        self.events(started["turn"])
        follow = self.client.post(f"/api/s5/threads/{started['thread']}/turns", headers=H,
                                  json={"question": "E rispetto a PostgreSQL?", "kind": "research"}).json()
        self.events(follow["turn"])
        self.assertEqual(self.fakes.briefs[-1]["question"], "E rispetto a PostgreSQL?")
        self.assertIn("SQLite", self.fakes.briefs[-1]["previous"])
        self.assertIn("PostgreSQL", self.fakes.calls[-1]["task"])

    def test_brief_failure_falls_back_to_question(self):
        self.fakes.brief_fails = True
        started = self.start()
        done = self.events(started["turn"])[-1]["turn"]
        self.assertEqual(done["status"], "done")
        self.assertEqual(done["brief"], "")
        self.assertEqual(self.fakes.calls[-1]["task"], "Spiega in italiano quando usare la modalità WAL di SQLite")

    def test_cancel_marks_turn_cancelled(self):
        self.fakes.gate = asyncio.Event()
        started = self.start()
        response = self.client.post(f"/api/s5/turns/{started['turn']}/cancel", headers=H)
        self.assertTrue(response.json()["cancelled"])
        done = self.events(started["turn"])[-1]["turn"]
        self.assertEqual(done["status"], "cancelled")

    def test_delete_removes_thread_and_only_its_output_files(self):
        started = self.start()
        self.events(started["turn"])
        keep = Path("outputs/other.md")
        keep.write_text("x")
        response = self.client.delete(f"/api/s5/threads/{started['thread']}", headers=H)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Path("outputs/t1.md").exists())
        self.assertTrue(keep.exists())
        self.assertEqual(self.client.get(f"/api/s5/threads/{started['thread']}").status_code, 404)

    def test_validation_and_csrf_header(self):
        bad = [
            ({"question": ""}, H), ({"question": "x" * 2001}, H), ({"question": "ok", "mode": "turbo"}, H),
            ({"question": "ok", "max_results": 99}, H), ({"question": "ok", "max_results": "8"}, H),
            ({"question": "ok", "tone": "Evil"}, H), ({"question": "ok", "lang": "xx"}, H),
            ({"question": "ok", "ui": "de"}, H),
        ]
        for body, headers in bad:
            with self.subTest(body=body):
                self.assertEqual(self.client.post("/api/s5/threads", headers=headers, json=body).status_code, 400)
        self.assertEqual(self.client.post("/api/s5/threads", json={"question": "ok"}).status_code, 403)
        self.assertEqual(self.client.delete("/api/s5/threads/x").status_code, 403)
        self.assertEqual(self.client.post("/api/s5/threads/missing/turns", headers=H,
                                          json={"question": "ok"}).status_code, 404)
        self.assertEqual(self.client.get("/api/s5/turns/missing/events").status_code, 404)


class UnitTests(unittest.TestCase):
    def test_restart_fails_interrupted_turns(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(Path(tmp) / "s.db")
            thread = store.create_thread("q")
            turn = store.add_turn(thread, kind="research", question="q", mode="quick", lang="it",
                                  tone="Objective", max_results=8)
            self.assertEqual(Store(Path(tmp) / "s.db").fail_interrupted(), 1)
            self.assertEqual(store.turn(turn)["status"], "error")
            self.assertEqual(store.turn(turn)["error"], "interrupted")

    def test_existing_database_gains_brief_column(self):
        import sqlite3
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "old.db"
            old = sqlite3.connect(path)
            old.executescript(SCHEMA_BEFORE_BRIEF)
            old.execute("INSERT INTO threads VALUES ('t','q',1,1)")
            old.execute("INSERT INTO turns(id,thread_id,seq,kind,question,mode,lang,tone,max_results,status,created)"
                        " VALUES ('u','t',1,'research','q','quick','it','Objective',8,'done',1)")
            old.commit()
            old.close()
            store = Store(path)
            self.assertEqual(store.turn("u")["brief"], "")
            self.assertEqual((store.turn("u")["tokens_in"], store.turn("u")["llm_calls"]), (0, 0))
            store.update_turn("u", brief="b")
            self.assertEqual(store.turn("u")["brief"], "b")

    def test_concurrent_meters_do_not_mix(self):
        from star5.tokens import metered

        async def job(i, o):
            with metered() as meter:
                await fake_model_call(i, o)
                await asyncio.sleep(0.01)
                await fake_model_call(i, o)
                return meter.snapshot()

        async def both():
            return await asyncio.gather(job(10, 1), job(1000, 100))

        first, second = asyncio.run(both())
        self.assertEqual(first, {"tokens_in": 20, "tokens_out": 2, "llm_calls": 2})
        self.assertEqual(second, {"tokens_in": 2000, "tokens_out": 200, "llm_calls": 2})

    def test_normalize_ignores_images_and_unknown(self):
        self.assertEqual(normalize({"type": "images", "output": "[]"}), [])
        self.assertEqual(normalize("x"), [])
        self.assertEqual(normalize({"type": "logs", "content": "writing_report", "output": "w"}),
                         [{"t": "phase", "phase": "write"}, {"t": "log", "text": "w"}])

    def test_build_sources_dedupes_and_strips_punctuation(self):
        sources = build_sources("a https://a.org/x. b https://a.org/x", [], ["https://b.org"])
        self.assertEqual([s["url"] for s in sources], ["https://a.org/x", "https://b.org"])

    def test_detect_lang_fallback(self):
        self.assertEqual(detect_lang("ok", "fr"), "fr")
        self.assertEqual(detect_lang("Quali sono i vantaggi della modalità WAL in SQLite?", "en"), "it")

    def test_drop_event_logs_only_removes_this_runs_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            old = web.OUTPUTS
            web.OUTPUTS = Path(tmp)
            try:
                since = time.time()
                mine, other, report = (Path(tmp) / "task_1_abc.json", Path(tmp) / "task_1_zzz.json",
                                       Path(tmp) / "task_1_abc.md")
                for path in (mine, other, report):
                    path.write_text("x")
                earlier = Path(tmp) / "task_0_abc.json"
                earlier.write_text("x")
                os.utime(earlier, (since - 3600, since - 3600))
                drop_event_logs("abc", since)
                self.assertFalse(mine.exists())
                self.assertTrue(other.exists() and report.exists() and earlier.exists())
            finally:
                web.OUTPUTS = old


if __name__ == "__main__":
    unittest.main()
